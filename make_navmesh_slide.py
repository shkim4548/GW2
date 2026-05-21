# -*- coding: utf-8 -*-
"""NAVMESH_EXPORTER_SLIDE.pptx — NavMesh / Map 데이터 익스포터 한페이지 (v3 라이트 테마)"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

W, H = Inches(13.33), Inches(7.5)

WHITE   = RGBColor(0xFF,0xFF,0xFF)
ACCENT  = RGBColor(0x2B,0x5B,0xA8)
DARK    = RGBColor(0x1A,0x1A,0x2E)
GRAY    = RGBColor(0x6B,0x72,0x80)
LGRAY   = RGBColor(0xD0,0xD5,0xDD)
BLU_BG  = RGBColor(0xEA,0xF3,0xFF); BLU_BD = RGBColor(0x44,0x72,0xC4)
GRN_BG  = RGBColor(0xEA,0xFF,0xF0); GRN_BD = RGBColor(0x70,0xAD,0x47)
ORG_BG  = RGBColor(0xFF,0xF0,0xE0); ORG_BD = RGBColor(0xE6,0x7E,0x22)
YEL_BG  = RGBColor(0xFF,0xFB,0xCC); YEL_BD = RGBColor(0xF0,0xC0,0x30)
T_HDR   = RGBColor(0xE8,0xEC,0xEF)
T_BLU   = RGBColor(0xDB,0xEA,0xFF)
T_GRN   = RGBColor(0xDB,0xFF,0xEC)
T_YEL   = RGBColor(0xFF,0xFB,0xCC)
T_ORG   = RGBColor(0xFF,0xED,0xD5)
T_RED   = RGBColor(0xFF,0xE8,0xE8)
PNT_BG  = RGBColor(0xEF,0xF4,0xFF)
PNT_BD  = RGBColor(0x44,0x72,0xC4)
PNT_LB  = RGBColor(0xE6,0x7E,0x22)
CODE_BG = RGBColor(0xF6,0xF8,0xFA)
KR = "맑은 고딕"; CO = "Consolas"

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H
blank = prs.slide_layouts[6]

# ── 유틸 ──────────────────────────────────────────────────────────────────────
def rect(slide, x, y, w, h, fill=None, line=None, lw=Pt(1)):
    sh = slide.shapes.add_shape(1, x, y, w, h)
    sh.fill.background()
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line: sh.line.color.rgb = line; sh.line.width = lw
    else: sh.line.fill.background()
    return sh

def tb(slide, text, x, y, w, h, fn=KR, sz=Pt(11), bold=False, color=DARK,
       align=PP_ALIGN.LEFT, wrap=True):
    t = slide.shapes.add_textbox(x, y, w, h)
    tf = t.text_frame; tf.word_wrap = wrap
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = fn; r.font.size = sz
    r.font.bold = bold; r.font.color.rgb = color
    return t

def header(slide, title, subtitle=''):
    rect(slide, 0, 0, W, H, fill=WHITE)
    rect(slide, Inches(0.35), Inches(0.28), Inches(0.07), Inches(0.90), fill=ACCENT)
    tb(slide, title, Inches(0.52), Inches(0.22), Inches(12.3), Inches(0.55),
       sz=Pt(30), bold=True, color=DARK)
    if subtitle:
        tb(slide, subtitle, Inches(0.52), Inches(0.78), Inches(12.3), Inches(0.30),
           sz=Pt(12), color=GRAY)
    ln = slide.shapes.add_connector(1, Inches(0.35), Inches(1.15), Inches(12.98), Inches(1.15))
    ln.line.color.rgb = LGRAY; ln.line.width = Pt(0.75)

def point_box(slide, text, y=Inches(6.62)):
    rect(slide, Inches(0.35), y, Inches(12.63), Inches(0.65), fill=PNT_BG, line=PNT_BD)
    rect(slide, Inches(0.35), y, Inches(1.05),  Inches(0.65), fill=PNT_LB)
    tb(slide, 'POINT', Inches(0.37), y+Inches(0.13), Inches(1.0), Inches(0.40),
       sz=Pt(11), bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    tb(slide, text, Inches(1.50), y+Inches(0.07), Inches(11.3), Inches(0.52),
       sz=Pt(10.5), color=DARK)

def step_box(slide, num, title_txt, body, x, y, w, h, bg, bd):
    rect(slide, x,             y, Inches(0.06), h, fill=bd)
    rect(slide, x+Inches(0.06),y, w-Inches(0.06), h, fill=bg, line=bd, lw=Pt(0.5))
    tb(slide, f'{num}  {title_txt}',
       x+Inches(0.14), y+Inches(0.05), w-Inches(0.20), Inches(0.26),
       sz=Pt(10.5), bold=True, color=bd)
    tb(slide, body,
       x+Inches(0.14), y+Inches(0.29), w-Inches(0.20), h-Inches(0.34),
       sz=Pt(9.5), color=DARK)

def cbox(slide, body, x, y, w, h, bg, bd, title='', tsz=Pt(10.5)):
    rect(slide, x, y, w, h, fill=bg, line=bd, lw=Pt(1.2))
    if title:
        tb(slide, title, x+Inches(0.1), y+Inches(0.06), w-Inches(0.2), Inches(0.28),
           sz=tsz, bold=True, color=bd)
        tb(slide, body,  x+Inches(0.1), y+Inches(0.32), w-Inches(0.2), h-Inches(0.38),
           sz=Pt(10), color=DARK)
    else:
        tb(slide, body, x+Inches(0.1), y+Inches(0.1), w-Inches(0.2), h-Inches(0.2),
           sz=Pt(10), color=DARK)

def table(slide, headers, rows, row_colors, x, y, col_widths, row_h=Inches(0.44)):
    nc = len(headers); nr = len(rows)+1
    tbl = slide.shapes.add_table(nr, nc, x, y, sum(col_widths), row_h*nr).table
    for ci, cw in enumerate(col_widths): tbl.columns[ci].width = cw

    def cfill(cell, color):
        tc = cell._tc; tcPr = tc.get_or_add_tcPr()
        sf = parse_xml(f'<a:solidFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                       f'<a:srgbClr val="{str(color)}"/></a:solidFill>')
        for old in tcPr.findall(qn('a:solidFill')): tcPr.remove(old)
        tcPr.insert(0, sf)

    def ctext(cell, text, bold=False, sz=Pt(10), al=PP_ALIGN.LEFT):
        cell.text = ''
        tf = cell.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = al
        r = p.add_run(); r.text = text
        r.font.name = KR; r.font.size = sz
        r.font.bold = bold; r.font.color.rgb = DARK

    for ci, h in enumerate(headers):
        c = tbl.cell(0, ci); cfill(c, T_HDR)
        ctext(c, h, bold=True, sz=Pt(10.5), al=PP_ALIGN.CENTER)
    for ri, (row, rc) in enumerate(zip(rows, row_colors)):
        for ci, val in enumerate(row):
            c = tbl.cell(ri+1, ci); cfill(c, rc); ctext(c, val)
    for ri in range(nr): tbl.rows[ri].height = row_h
    return tbl

# ══════════════════════════════════════════════════════════════════════════════
#  슬라이드: NavMesh / Map 데이터 익스포터
# ══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
header(s, 'NavMesh / Map 데이터 익스포터',
       'GW2_Client  Assets/Editor  —  Unity 씬 데이터를 C++ 서버용 파일로 직접 변환하는 Editor Tool 파이프라인')

# ── 왼쪽: 4개 도구 스텝 박스 ────────────────────────────────────────────────
LX = Inches(0.35)
LW = Inches(6.10)
SH = Inches(0.96)
GAP = Inches(0.06)
Y0  = Inches(1.28)

tools = [
    ('①', 'NavmeshJsonExporter  —  삼각형 메시 → JSON',
     'NavMesh.CalculateTriangulation()  →  vertices[][3] + indices[] 배열 직렬화\n'
     '출력: Assets/NavMeshExport/navmesh.json  (메뉴: Tools/Export/Export NavMesh to JSON)',
     BLU_BG, BLU_BD),

    ('②', 'NavGridBinExporter  —  격자 보행가능 맵 → BIN',
     'cellSize(기본 0.5) 단위로 맵 전체 격자 분할  →  각 셀 중심에 Raycast + NavMesh.SamplePosition\n'
     '출력: navgrid.bin  (MAGIC NRGD  헤더 + walkable byte[] 배열)',
     GRN_BG, GRN_BD),

    ('③', 'LaneMapExporter  —  Tilemap 레인 구역 → BIN',
     'Top / Mid / Bot Tilemap 오브젝트 → 셀별 laneId(1/2/3) 매핑  (navgrid.bin 스펙 공유)\n'
     '출력: laneMap.bin  (MAGIC LNMP  + byte[W×H])',
     YEL_BG, YEL_BD),

    ('④', 'LaneRouteExporter  —  웨이포인트 경로 → JSON',
     'LaneWaypoint 씬 오브젝트 수집 (WP_0, WP_1… 이름 순 정렬)  →  laneId + waypoints[] 직렬화\n'
     '출력: laneRoutes.json  (메뉴: Tools/Lane/Lane Route Exporter)',
     ORG_BG, ORG_BD),
]
for idx, (num, ttxt, body, bg, bd) in enumerate(tools):
    step_box(s, num, ttxt, body, LX, Y0 + idx*(SH+GAP), LW, SH, bg, bd)

# 흐름 요약 박스
FY = Y0 + 4*(SH+GAP)
cbox(s,
     'Unity NavMesh Bake  →  Editor Tool 실행  →  Assets/NavMeshExport/  →  C++ 서버 로드',
     LX, FY, LW, Inches(0.62),
     CODE_BG, LGRAY)

# ── 오른쪽 ────────────────────────────────────────────────────────────────────
RX = Inches(6.62)
RW = Inches(6.36)

# navgrid.bin 바이너리 포맷 박스
cbox(s,
     'MAGIC   uint32   0x4E524744  ("NRGD")\n'
     'VERSION uint32   1\n'
     'width   int32    격자 X 셀 수\n'
     'height  int32    격자 Z 셀 수\n'
     'cellSize float   셀 크기 (기본 0.5 m)\n'
     'origin  float[3] 격자 시작 좌표 (x, y, z)\n'
     'data    byte[]   walkable 플래그  (1=가능, 0=불가)',
     RX, Inches(1.28), RW, Inches(1.78),
     CODE_BG, GRN_BD, title='navgrid.bin  바이너리 레이아웃', tsz=Pt(10.5))

# 4개 도구 비교 표
tb(s, '도구별 출력 파일 요약', RX, Inches(3.15), RW, Inches(0.26),
   sz=Pt(11), bold=True, color=DARK)

tbl_rows = [
    ['NavmeshJsonExporter', 'navmesh.json',   '삼각형 정점/인덱스',     '서버 충돌·가시성 계산'],
    ['NavGridBinExporter',  'navgrid.bin',    'walkable 격자 맵',       '서버 A* / 경로 탐색'],
    ['LaneMapExporter',     'laneMap.bin',    '레인별 셀 ID 배열',      '미니언 소환 구역 판단'],
    ['LaneRouteExporter',   'laneRoutes.json','레인 웨이포인트 목록',   '미니언 이동 경로 지정'],
]
table(s, ['도구', '출력 파일', '포함 데이터', '서버 활용'],
      tbl_rows, [T_BLU, T_GRN, T_YEL, T_ORG],
      RX, Inches(3.48),
      [Inches(1.90), Inches(1.30), Inches(1.60), Inches(1.56)],
      row_h=Inches(0.52))

# Bounds 결정 로직 박스
cbox(s,
     '① Manual Bounds 지정  →  직접 Min/Max 입력\n'
     '② Map Root GameObject  →  자식 Renderer 전체 Encapsulate\n'
     '③ NavMesh 자동  →  CalculateTriangulation() 정점 범위 (기본값)',
     RX, Inches(5.62), RW, Inches(0.76),
     YEL_BG, YEL_BD, title='NavGridBinExporter  Bounds 결정 우선순위')

point_box(s,
    'Unity NavMesh.CalculateTriangulation() 로 얻은 삼각형 데이터를 서버가 직접 읽을 수 있는 '
    'JSON/BIN 포맷으로 변환. 격자(Grid) 방식으로 이중 저장해 A* 탐색 성능 최적화. '
    'laneMap + laneRoute 로 미니언 소환·이동 로직을 서버에서 완전 구현')

# ── 저장 ──────────────────────────────────────────────────────────────────────
OUT = r'D:\Dev\unity\GW2\NAVMESH_EXPORTER_SLIDE.pptx'
prs.save(OUT)
print(f'저장 완료: {OUT}')
print(f'슬라이드 수: {len(prs.slides)}')
