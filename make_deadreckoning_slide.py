# -*- coding: utf-8 -*-
"""DEAD_RECKONING_SLIDE.pptx — 클라이언트 예측 & 추측항법 한페이지 (v3 라이트 테마)"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

W, H = Inches(13.33), Inches(7.5)

WHITE   = RGBColor(0xFF,0xFF,0xFF)
ACCENT  = RGBColor(0x2B,0x5B,0xA8)
DARK    = RGBColor(0x1A,0x1A,0x2E)
GRAY    = RGBColor(0x6B,0x72,0x80)
LGRAY   = RGBColor(0xD0,0xD5,0xDD)
BLU_BG  = RGBColor(0xEA,0xF3,0xFF); BLU_BD = RGBColor(0x44,0x72,0xC4)
GRN_BG  = RGBColor(0xEA,0xFF,0xF0); GRN_BD = RGBColor(0x70,0xAD,0x47)
ORG_BG  = RGBColor(0xFF,0xF0,0xE0); ORG_BD = RGBColor(0xE6,0x7E,0x22)
YEL_BG  = RGBColor(0xFF,0xFB,0xCC); YEL_BD = RGBColor(0xF0,0xC0,0x30)
T_HDR   = RGBColor(0xE8,0xEC,0xEF)
T_BLU   = RGBColor(0xDB,0xEA,0xFF)
T_GRN   = RGBColor(0xDB,0xFF,0xEC)
T_YEL   = RGBColor(0xFF,0xFB,0xCC)
T_ORG   = RGBColor(0xFF,0xED,0xD5)
PNT_BG  = RGBColor(0xEF,0xF4,0xFF)
PNT_BD  = RGBColor(0x44,0x72,0xC4)
PNT_LB  = RGBColor(0xE6,0x7E,0x22)
CODE_BG = RGBColor(0xF6,0xF8,0xFA)
KR = "맑은 고딕"; CO = "Consolas"

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H
blank = prs.slide_layouts[6]

# ── 유틸 ──────────────────────────────────────────────────────────────────────
def rect(slide, x, y, w, h, fill=None, line=None, lw=Pt(1)):
    sh = slide.shapes.add_shape(1, x, y, w, h)
    sh.fill.background()
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line: sh.line.color.rgb = line; sh.line.width = lw
    else: sh.line.fill.background()
    return sh

def tb(slide, text, x, y, w, h, fn=KR, sz=Pt(11), bold=False, color=DARK,
       align=PP_ALIGN.LEFT, wrap=True):
    t = slide.shapes.add_textbox(x, y, w, h)
    tf = t.text_frame; tf.word_wrap = wrap
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = fn; r.font.size = sz
    r.font.bold = bold; r.font.color.rgb = color
    return t

def header(slide, title, subtitle=''):
    rect(slide, 0, 0, W, H, fill=WHITE)
    rect(slide, Inches(0.35), Inches(0.28), Inches(0.07), Inches(0.90), fill=ACCENT)
    tb(slide, title, Inches(0.52), Inches(0.22), Inches(12.3), Inches(0.55),
       sz=Pt(30), bold=True, color=DARK)
    if subtitle:
        tb(slide, subtitle, Inches(0.52), Inches(0.78), Inches(12.3), Inches(0.30),
           sz=Pt(12), color=GRAY)
    ln = slide.shapes.add_connector(1, Inches(0.35), Inches(1.15), Inches(12.98), Inches(1.15))
    ln.line.color.rgb = LGRAY; ln.line.width = Pt(0.75)

def point_box(slide, text, y=Inches(6.62)):
    rect(slide, Inches(0.35), y, Inches(12.63), Inches(0.65), fill=PNT_BG, line=PNT_BD)
    rect(slide, Inches(0.35), y, Inches(1.05),  Inches(0.65), fill=PNT_LB)
    tb(slide, 'POINT', Inches(0.37), y+Inches(0.13), Inches(1.0), Inches(0.40),
       sz=Pt(11), bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    tb(slide, text, Inches(1.50), y+Inches(0.07), Inches(11.3), Inches(0.52),
       sz=Pt(10.5), color=DARK)

def step_box(slide, num, title_txt, body, x, y, w, h, bg, bd):
    rect(slide, x,             y, Inches(0.06), h, fill=bd)
    rect(slide, x+Inches(0.06),y, w-Inches(0.06), h, fill=bg, line=bd, lw=Pt(0.5))
    tb(slide, f'{num}  {title_txt}',
       x+Inches(0.14), y+Inches(0.05), w-Inches(0.20), Inches(0.26),
       sz=Pt(10.5), bold=True, color=bd)
    tb(slide, body,
       x+Inches(0.14), y+Inches(0.29), w-Inches(0.20), h-Inches(0.34),
       sz=Pt(9.5), color=DARK)

def cbox(slide, body, x, y, w, h, bg, bd, title='', tsz=Pt(10.5)):
    rect(slide, x, y, w, h, fill=bg, line=bd, lw=Pt(1.2))
    if title:
        tb(slide, title, x+Inches(0.1), y+Inches(0.06), w-Inches(0.2), Inches(0.28),
           sz=tsz, bold=True, color=bd)
        tb(slide, body,  x+Inches(0.1), y+Inches(0.32), w-Inches(0.2), h-Inches(0.38),
           sz=Pt(10), color=DARK)
    else:
        tb(slide, body, x+Inches(0.1), y+Inches(0.1), w-Inches(0.2), h-Inches(0.2),
           sz=Pt(10), color=DARK)

def table(slide, headers, rows, row_colors, x, y, col_widths, row_h=Inches(0.44)):
    nc = len(headers); nr = len(rows)+1
    tbl = slide.shapes.add_table(nr, nc, x, y, sum(col_widths), row_h*nr).table
    for ci, cw in enumerate(col_widths): tbl.columns[ci].width = cw

    def cfill(cell, color):
        tc = cell._tc; tcPr = tc.get_or_add_tcPr()
        sf = parse_xml(f'<a:solidFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                       f'<a:srgbClr val="{str(color)}"/></a:solidFill>')
        for old in tcPr.findall(qn('a:solidFill')): tcPr.remove(old)
        tcPr.insert(0, sf)

    def ctext(cell, text, bold=False, sz=Pt(10), al=PP_ALIGN.LEFT):
        cell.text = ''
        tf = cell.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = al
        r = p.add_run(); r.text = text
        r.font.name = KR; r.font.size = sz
        r.font.bold = bold; r.font.color.rgb = DARK

    for ci, h in enumerate(headers):
        c = tbl.cell(0, ci); cfill(c, T_HDR)
        ctext(c, h, bold=True, sz=Pt(10.5), al=PP_ALIGN.CENTER)
    for ri, (row, rc) in enumerate(zip(rows, row_colors)):
        for ci, val in enumerate(row):
            c = tbl.cell(ri+1, ci); cfill(c, rc); ctext(c, val)
    for ri in range(nr): tbl.rows[ri].height = row_h
    return tbl

# ══════════════════════════════════════════════════════════════════════════════
#  슬라이드: 클라이언트 예측 & 추측항법 (Dead Reckoning)
# ══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
header(s, '클라이언트 예측 & 추측항법 (Dead Reckoning)',
       'GW2 C++ 서버  ↔  Unity 클라이언트  —  100ms 브로드캐스트 기반 위치 동기화 전략')

# ── 왼쪽: 4-Step 보정 전략 ────────────────────────────────────────────────────
LX = Inches(0.35)
LW = Inches(6.10)
SH = Inches(0.86)
GAP = Inches(0.06)
Y0  = Inches(1.28)

steps = [
    ('①', '클라이언트 예측  —  내 플레이어 (MyPlayerController)',
     '클릭 즉시 NavMesh.CalculatePath()로 로컬 경로 계산 → 서버 응답 전 즉각 이동 시작\n'
     'CorrectPosition()  Lerp(speed=1.0)  ·  오차 < 0.01f → 즉시 Snap  ·  NavMeshAgent.updatePosition=false',
     BLU_BG, BLU_BD),
    ('②', '위치 보간  —  원격 플레이어 (PlayerController)',
     'S_MOVE(목적지 PosInfo, Yaw) 수신 → MoveTowards(_remoteSpeed=12.0)로 매 프레임 추종\n'
     'Slerp(speed=10.0) 회전 보간  ·  GetGroundY() NavMesh Y 보정으로 지형 고저차 처리',
     GRN_BG, GRN_BD),
    ('③', '스냅 보정  —  미니언 (MinionController)',
     '서버 제공 navPath 웨이포인트 순서대로 이동  ·  이동 중 Lerp 완전 비활성화 (진동 방지)\n'
     '오차 ≥ 5.0f → 즉시 Snap + FindNearestForwardWaypointIndex()로 경로 역행 방지',
     YEL_BG, YEL_BD),
    ('④', '서버 브로드캐스트  —  100ms / 10Hz  (Object.cpp)',
     'MOVE_BROADCAST_INTERVAL=0.1f  ·  C_MOVE 수신 시 MarkForceBroadcastMove() → 즉시 전파\n'
     'S_MOVE: 목적지 좌표 + Yaw 전송 (현재 위치 X)  ·  S_MOVE_END: 이동 완료 시 최종 위치 확정',
     ORG_BG, ORG_BD),
]
for idx, (num, ttxt, body, bg, bd) in enumerate(steps):
    step_box(s, num, ttxt, body, LX, Y0 + idx*(SH+GAP), LW, SH, bg, bd)

# 흐름 요약 박스
FY = Y0 + 4*(SH+GAP)
cbox(s,
     '클릭  →  StartMovePrediction() + C_MOVE(start, target, clientTime)  →  서버 pathfind  '
     '→  S_MOVE 100ms 브로드캐스트  →  CorrectPosition() / InterpolateToServerPosition()',
     LX, FY, LW, Inches(0.66), CODE_BG, LGRAY)

# ── 오른쪽 ────────────────────────────────────────────────────────────────────
RX = Inches(6.62)
RW = Inches(6.36)

# 보정 파라미터 비교 표
tb(s, '엔티티별 보정 전략 비교', RX, Inches(1.28), RW, Inches(0.26),
   sz=Pt(11), bold=True, color=DARK)

tbl_rows = [
    ['내 플레이어',   'NavMesh 예측\n+ Lerp 보정',    'Lerp speed=1.0\nSnap < 0.01f',    'NavMeshAgent\nupdatePosition=false'],
    ['원격 플레이어', 'MoveTowards\n+ Slerp 보간',   'speed=12.0\nrot speed=10.0',       'NavMesh Y 보정\nGetGroundY()'],
    ['미니언',        '서버 navPath\n+ Snap 보정',    'Snap ≥ 5.0f\nthreshold=2.5f',     '이동 중 Lerp\n비활성화'],
]
table(s, ['대상', '방식', '임계값', '특이사항'],
      tbl_rows, [T_BLU, T_GRN, T_YEL],
      RX, Inches(1.60),
      [Inches(1.40), Inches(1.60), Inches(1.50), Inches(1.86)],
      row_h=Inches(0.58))

# 패킷 구조 박스
cbox(s,
     'C_MOVE  →  StartPos(x,y,z)  +  TargetPos(x,y,z)  +  ClientTime(ms)\n'
     'S_MOVE  →  PosInfo(목적지 x,z)  +  Yaw  +  MoveState(RUN/IDLE)\n'
     'S_MOVE_END  →  PosInfo(최종 확정 x,y,z)  +  MoveState',
     RX, Inches(3.82), RW, Inches(1.00),
     CODE_BG, BLU_BD, title='패킷 구조 (C_MOVE / S_MOVE / S_MOVE_END)')

# 공격 범위 lag 보상 박스
cbox(s,
     'effectiveRange = attackRange + 2.0f  (스킬 ID=1 기준)\n'
     '근거: 이동속도 12 unit/s × RTT 50ms ≈ 0.6 unit  →  안전 마진 포함 +2.0f\n'
     '클라이언트에서 범위 내로 보여도 서버에서 범위 밖일 수 있는 lag 오차 보정',
     RX, Inches(4.92), RW, Inches(1.00),
     YEL_BG, YEL_BD, title='공격 범위 Lag 보상 (ClientPacketHandler)')

point_box(s,
    '서버는 목적지만 전송(현재 위치 X) → 클라이언트가 NavMesh 예측으로 입력 즉응. '
    '미니언은 이동 중 Lerp 비활성화로 진동 제거, 플레이어는 Lerp로 부드러운 보정. '
    '100ms 주기 + 강제 브로드캐스트(C_MOVE 수신 즉시)로 초기 동기화 보장')

# ── 저장 ──────────────────────────────────────────────────────────────────────
OUT = r'D:\Dev\unity\GW2\DEAD_RECKONING_SLIDE.pptx'
prs.save(OUT)
print(f'저장 완료: {OUT}')
print(f'슬라이드 수: {len(prs.slides)}')
